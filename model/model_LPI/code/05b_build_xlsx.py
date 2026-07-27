#!/usr/bin/env python3
"""
05b_build_xlsx.py — assemble the canonical LPI workbook from the CSV parts
written by 05_export_results.R.

Input : outputs/_xlsx_parts/*.csv   (NN_sheetname.csv; NN sets sheet order)
Output: outputs/lpi_scores.xlsx     (one sheet per part, styled header row)

Run from: model/model_LPI/  (after 05)   e.g.  .venv-python code/05b_build_xlsx.py
Uses openpyxl (already available in the project .venv).
"""
import csv, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PARTS = "outputs/_xlsx_parts"
OUT   = "outputs/lpi_scores.xlsx"

HDR_FILL = PatternFill("solid", fgColor="2B4A6F")
HDR_FONT = Font(bold=True, color="FFFFFF")

def sheet_name(fn):
    base = os.path.splitext(os.path.basename(fn))[0]
    return base.split("_", 1)[1][:31]  # strip NN_ order prefix, Excel 31-char cap

def is_num(s):
    try:
        float(s); return True
    except (ValueError, TypeError):
        return False

wb = Workbook(); wb.remove(wb.active)
for fn in sorted(glob.glob(f"{PARTS}/*.csv")):
    ws = wb.create_sheet(sheet_name(fn))
    with open(fn, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(rows[r-1], start=1):
            cell = ws.cell(row=r, column=c,
                           value=float(val) if (r > 1 and is_num(val)) else val)
            if r == 1:
                cell.fill = HDR_FILL; cell.font = HDR_FONT
                cell.alignment = Alignment(horizontal="left")
    ws.freeze_panes = "A2"
    # column widths from max content length
    for c in range(1, len(rows[0]) + 1):
        w = max((len(str(rows[r][c-1])) for r in range(len(rows)) if c-1 < len(rows[r])), default=10)
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 48)

wb.save(OUT)
print(f"wrote {OUT} with sheets: {wb.sheetnames}")
