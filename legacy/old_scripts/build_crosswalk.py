"""
build_crosswalk.py — complete source-to-final column crosswalk for the PHK pipeline.
Encodes the user's domain-based naming spec verbatim, matches it to the ACTUAL
workbook columns, flags anything uncovered, and writes:
  data/processed/indicator_dictionary.csv   (full mapping + dictionary)
Run from repo root:  python code/build_crosswalk.py
"""
import csv, re, os
from openpyxl import load_workbook
from naming import shorten

PHKX='data/raw_data/Data untuk PHK Dashboard.xlsx'
MAPX='data/raw_data/MAP_composite_all.xlsx'
os.makedirs('data/processed', exist_ok=True)

def norm(s):
    return re.sub(r'\s+',' ', str(s).replace('\n',' ')).strip()

# ---- verbatim user mappings (orig -> final), per sheet ----
BULAN = """
Provinsi → province_name
Tahun → year
Bulan → month
Periode → date
PHK (Stock) → phk_stock
PHK (Flow) → phk_flow
PMI Manufaktur S&P → macro_pmi_manufaktur_national
IHK (2010=100) → price_cpi_index
Inflasi Bulanan (M-to-M) → price_inflation_mom_pct
Inflasi Tahunan (Y-on-Y) → price_inflation_yoy_pct
Perubahan Harga Konsumen → price_consumer_change_pct
Nilai Ekspor → trade_export_value
Nilai Impor → trade_import_value
Neraca Perdagangan → trade_balance
NPL (IDR miliar) → fin_npl_idr_billion
BI Rate → macro_bi_rate_pct_national
Kurs → macro_exchange_rate_idr_usd_national
Harga Minyak Brent (USD/barrel) → price_brent_oil_usd_per_barrel_national
IHPB → price_ihpb_national
"""
TRIWULAN = """
Provinsi → province_name
Tahun → year
Triwulan → quarter
Bulan → quarter_end_month_label
PHK (Stock) → phk_stock_quarterly
PHK (Flow) → phk_flow_quarterly
PDRB (IDR mn) → macro_pdrb_idr_million_quarterly
PDRB Manufaktur (IDR mn) → macro_pdrb_manufacturing_idr_million_quarterly
PDRB Agrikultur (IDR mn) → macro_pdrb_agriculture_idr_million_quarterly
PDRB Jasa → macro_pdrb_services_idr_million_quarterly
Manufacturing Share (% PDRB) → macro_pdrb_manufacturing_share_pct_quarterly
Agriculture Share (% PDRB) → macro_pdrb_agriculture_share_pct_quarterly
Services Share (% PDRB) → macro_pdrb_services_share_pct_quarterly
PMI Manufaktur → macro_pmi_manufaktur_national_quarterly
FDI (USD mn) → fin_fdi_usd_million_quarterly
FDI (IDR mn) → fin_fdi_idr_million_quarterly
Jumlah Penduduk Bekerja → lab_working_population_quarterly
IHP (2016=100) → price_producer_index_quarterly
Perubahan Harga Produsen** → price_producer_change_pct_quarterly
Perubahan Harga Produsen (CEIC) → price_producer_change_ceic_pct_quarterly
IHK (2010=100) → price_cpi_index_quarterly
Perubahan Harga Konsumen → price_consumer_change_pct_quarterly
Nilai Ekspor (USD juta) → trade_export_value_usd_million_quarterly
Nilai Impor (USD juta) → trade_import_value_usd_million_quarterly
Neraca Perdagangan* → trade_balance_usd_million_quarterly
NPL, End of Quarter (IDR miliar) → fin_npl_idr_billion_quarterly
BI Rate → macro_bi_rate_pct_national_quarterly
Kurs → macro_exchange_rate_idr_usd_national_quarterly
IKK → macro_consumer_confidence_index_quarterly
"""
TAHUN = """
Provinsi → province_name
Tahun → year
PHK → phk_annual
% Pekerja Formal → lab_formal_share_pct_annual
% Pekerja Formal (2019) → lab_formal_share_2019_pct_annual
% Pekerja Formal (Growth from 2019, pp) → lab_formal_share_growth_from_2019_pp_annual
Jumlah Pekerja Memiliki Kontrak Tertulis → lab_written_contract_workers_annual
Jumlah Pekerja Formal → lab_formal_workers_annual
% Pekerja Memiliki Kontrak Tertulis → lab_written_contract_share_pct_annual
% Pekerja Memiliki Kontrak Tertulis (2019) → lab_written_contract_share_2019_pct_annual
% Pekerja Memiliki Kontrak Tertulis (Growth from 2019, pp) → lab_written_contract_share_growth_from_2019_pp_annual
Jumlah Penduduk Bekerja → lab_working_population_annual
% TPT → lab_tpt_pct_annual
% TPT Laki-laki → lab_tpt_male_pct_annual
% TPT Perempuan → lab_tpt_female_pct_annual
% TPAK → lab_tpak_pct_annual
% TPAK Laki-laki → lab_tpak_male_pct_annual
% TPAK Perempuan → lab_tpak_female_pct_annual
Berusaha Sendiri → lab_self_employed_annual
Buruh/karyawan/pegawai → lab_employee_count_annual
Pekerja Keluarga Tidak Dibayar → lab_unpaid_family_workers_annual
Penduduk Bekerja Setengah Penganggur → lab_underemployed_workers_annual
% Pekerja Penuh Waktu → lab_full_time_share_pct_annual
% Pekerja Paruh Waktu → lab_part_time_share_pct_annual
% Setengah Pengangguran → lab_underemployment_share_pct_annual
Jam Kerja → lab_working_hours_annual
Proporsi Lapangan Kerja Informal Sektor Non-Pertanian → lab_nonagri_informal_employment_share_pct_annual
Proporsi Lapangan Kerja Informal → lab_informal_employment_share_pct_annual
Pencari Kerja Terdaftar - Jumlah → lab_job_seekers_registered_annual
Lowongan Kerja Terdaftar - Jumlah → lab_vacancies_registered_annual
Penempatan/Pemenuhan Tenaga Kerja - Jumlah → lab_placements_registered_annual
Tenaga Kerja - Pertanian, Kehutanan dan Perikanan → emp_agriculture_forestry_fishery_annual
Tenaga Kerja - Pertambangan dan Penggalian → emp_mining_quarrying_annual
Tenaga Kerja - Industri Pengolahan → emp_manufacturing_annual
Tenaga Kerja - Pengadaan Listrik dan Gas → emp_electricity_gas_annual
Tenaga Kerja - Pengadaan Air, Pengelolaan Sampah, Limbah dan Daur Ulang → emp_water_waste_management_annual
Tenaga Kerja - Konstruksi → emp_construction_annual
Tenaga Kerja - Perdagangan Besar dan Eceran, Reparasi Mobil dan Sepeda Motor → emp_trade_repair_annual
Tenaga Kerja - Transportasi dan Pergudangan → emp_transport_warehousing_annual
Tenaga Kerja - Penyediaan Akomodasi dan Makan Minum → emp_accommodation_food_annual
Tenaga Kerja - Informasi dan Komunikasi → emp_information_communication_annual
Tenaga Kerja - Jasa Keuangan dan Asuransi → emp_financial_insurance_annual
Tenaga Kerja - Real Estate → emp_real_estate_annual
Tenaga Kerja - Jasa Perusahaan → emp_business_services_annual
Tenaga Kerja - Administrasi Pemerintahan, Pertahanan dan Jaminan Sosial Wajib → emp_public_admin_defense_social_security_annual
Tenaga Kerja - Jasa Pendidikan → emp_education_annual
Tenaga Kerja - Jasa Kesehatan dan Kegiatan Sosial → emp_health_social_annual
Tenaga Kerja - Jasa Lainnya → emp_other_services_annual
% Share Tenaga Kerja - Pertanian, Kehutanan dan Perikanan → emp_share_agriculture_forestry_fishery_pct_annual
% Share Tenaga Kerja - Pertambangan dan Penggalian → emp_share_mining_quarrying_pct_annual
% Share Tenaga Kerja - Industri Pengolahan → emp_share_manufacturing_pct_annual
% Share Tenaga Kerja - Pengadaan Listrik dan Gas → emp_share_electricity_gas_pct_annual
% Share Tenaga Kerja - Pengadaan Air, Pengelolaan Sampah, Limbah dan Daur Ulang → emp_share_water_waste_management_pct_annual
% Share Tenaga Kerja - Konstruksi → emp_share_construction_pct_annual
% Share Tenaga Kerja - Perdagangan Besar dan Eceran, Reparasi Mobil dan Sepeda Motor → emp_share_trade_repair_pct_annual
% Share Tenaga Kerja - Transportasi dan Pergudangan → emp_share_transport_warehousing_pct_annual
% Share Tenaga Kerja - Penyediaan Akomodasi dan Makan Minum → emp_share_accommodation_food_pct_annual
% Share Tenaga Kerja - Informasi dan Komunikasi → emp_share_information_communication_pct_annual
% Share Tenaga Kerja - Jasa Keuangan dan Asuransi → emp_share_financial_insurance_pct_annual
% Share Tenaga Kerja - Real Estate → emp_share_real_estate_pct_annual
% Share Tenaga Kerja - Jasa Perusahaan → emp_share_business_services_pct_annual
% Share Tenaga Kerja - Administrasi Pemerintahan, Pertahanan dan Jaminan Sosial Wajib → emp_share_public_admin_defense_social_security_pct_annual
% Share Tenaga Kerja - Jasa Pendidikan → emp_share_education_pct_annual
% Share Tenaga Kerja - Jasa Kesehatan dan Kegiatan Sosial → emp_share_health_social_pct_annual
% Share Tenaga Kerja - Jasa Lainnya → emp_share_other_services_pct_annual
UMP (IDR) → wage_ump_idr_annual
Growth UMP → wage_ump_growth_pct_annual
Upah buruh/karyawan/pegawai (IDR) → wage_avg_employee_idr_annual
Kaitz Index → wage_kaitz_index_annual
BPJSTK - Peserta Aktif PU → bpjstk_active_pu_annual
BPJSTK - Peserta Aktif BPU → bpjstk_active_bpu_annual
BPJSTK - PHK → bpjstk_phk_annual
BPJSTK - JHT PHK → bpjstk_jht_phk_annual
BPJSTK - JKP PHK → bpjstk_jkp_phk_annual
Pertumbuhan Bekerja → growth_lab_working_population_annual
Pertumbuhan Buruh/Karyawan/Pegawai → growth_lab_employee_count_annual
Pertumbuhan Pekerja Keluarga Tidak Dibayar → growth_lab_unpaid_family_workers_annual
Pertumbuhan % Pekerja Formal → growth_lab_formal_share_pct_annual
Pertumbuhan % Pekerja Memiliki Kontrak Tertulis → growth_lab_written_contract_share_pct_annual
Pertumbuhan % Pekerja Penuh Waktu → growth_lab_full_time_share_pct_annual
Pertumbuhan % Pekerja Paruh Waktu → growth_lab_part_time_share_pct_annual
Pertumbuhan % Setengah Pengangguran → growth_lab_underemployment_share_pct_annual
Pertumbuhan Jam Kerja → growth_lab_working_hours_annual
Pertumbuhan TPT → growth_lab_tpt_pct_annual
Pertumbuhan TPT Laki-laki → growth_lab_tpt_male_pct_annual
Pertumbuhan TPT Perempuan → growth_lab_tpt_female_pct_annual
Pertumbuhan TPAK → growth_lab_tpak_pct_annual
Pertumbuhan TPAK Laki-laki → growth_lab_tpak_male_pct_annual
Pertumbuhan TPAK Perempuan → growth_lab_tpak_female_pct_annual
Pertumbuhan Upah → growth_wage_avg_employee_annual
Pertumbuhan BPJSTK - Peserta Aktif PU → growth_bpjstk_active_pu_annual
Pertumbuhan BPJSTK - Peserta Aktif BPU → growth_bpjstk_active_bpu_annual
Pertumbuhan BPJSTK - PHK → growth_bpjstk_phk_annual
Pertumbuhan BPJSTK - JHT PHK → growth_bpjstk_jht_phk_annual
Pertumbuhan BPJSTK - JKP PHK → growth_bpjstk_jkp_phk_annual
Lag_1 PHK → lag1_phk_annual
Lag_1 % Pekerja Formal → lag1_lab_formal_share_pct_annual
Lag_1 % Pekerja Memiliki Kontrak Tertulis → lag1_lab_written_contract_share_pct_annual
Lag_1 Growth UMP → lag1_wage_ump_growth_pct_annual
Lag_1 BPJSTK - Peserta Aktif PU → lag1_bpjstk_active_pu_annual
Lag_1 BPJSTK - Peserta Aktif BPU → lag1_bpjstk_active_bpu_annual
Lag_1 BPJSTK - PHK → lag1_bpjstk_phk_annual
Lag_1 BPJSTK - JHT PHK → lag1_bpjstk_jht_phk_annual
Lag_1 BPJSTK - JKP PHK → lag1_bpjstk_jkp_phk_annual
Perusahaan Industri Sedang → ind_firms_medium_annual
Perusahaan Industri Besar → ind_firms_large_annual
Jumlah Perusahaan → ind_firms_total_annual
Tenaga Kerja Industri Sedang → ind_workers_medium_annual
Tenaga Kerja Industri Besar → ind_workers_large_annual
Jumlah Tenaga Kerja Produksi → ind_workers_production_annual
Jumlah Tenaga Kerja Lainnya → ind_workers_nonproduction_annual
Jumlah Tenaga Kerja → ind_workers_total_annual
Pengeluaran untuk Pekerja Produksi Upah/Gaji, Upah Lembur, Tunjangan → ind_production_workers_wage_cost_annual
Pengeluaran untuk Pekerja Produksi Lainnya → ind_production_workers_other_cost_annual
Jumlah Pengeluaran untuk Pekerja Produksi → ind_production_workers_total_cost_annual
Pengeluaran untuk Pekerja Lainnya Upah/Gaji, Upah Lembur, Tunjangan → ind_nonproduction_workers_wage_cost_annual
Pengeluaran untuk Pekerja Lainnya Lainnya → ind_nonproduction_workers_other_cost_annual
Jumlah Pengeluaran untuk Pekerja Lainnya → ind_nonproduction_workers_total_cost_annual
Jumlah Pengeluaran Seluruh Pekerja → ind_labor_cost_total_annual
Biaya Input → ind_input_cost_annual
Nilai Output → ind_output_value_annual
Nilai Tambah (harga pasar) → ind_value_added_market_annual
Nilai Tambah (biaya faktor produksi) → ind_value_added_factor_cost_annual
Jumlah Perusahaan PMDN → ind_firms_pmdn_annual
Jumlah Perusahaan PMA → ind_firms_pma_annual
Jumlah Tenaga Kerja/Jumlah Perusahaan → ind_workers_per_firm_annual
Jumlah Tenaga Kerja Produksi/Jumlah Tenaga Kerja → ind_production_worker_share_annual
Jumlah Tenaga Kerja/Biaya Input → ind_workers_per_input_cost_annual
Jumlah Tenaga Kerja/Nilai Output → ind_workers_per_output_value_annual
Jumlah Tenaga Kerja/Nilai Tambah (harga pasar) → ind_workers_per_value_added_market_annual
Jumlah Tenaga Kerja/Nilai Tambah (biaya faktor produksi) → ind_workers_per_value_added_factor_cost_annual
Biaya Pekerja/Biaya Input → ind_labor_cost_share_input_cost_annual
Biaya Pekerja/Nilai Output → ind_labor_cost_share_output_value_annual
Biaya Pekerja/Nilai Tambah Pasar → ind_labor_cost_share_value_added_market_annual
Biaya Pekerja/Nilai Tambah Produksi → ind_labor_cost_share_value_added_factor_cost_annual
Biaya Pekerja/Tenaga Kerja → ind_labor_cost_per_worker_annual
Nilai Output/Tenaga Kerja → ind_output_value_per_worker_annual
Nilai Tambah Pasar/Tenaga Kerja → ind_value_added_market_per_worker_annual
Nilai Tambah Produksi/Tenaga Kerja → ind_value_added_factor_cost_per_worker_annual
Biaya Input/Nilai Output → ind_input_cost_share_output_value_annual
Nilai Tambah Pasar/Nilai Output → ind_value_added_market_share_output_value_annual
Nilai Tambah Produksi/Nilai Output → ind_value_added_factor_cost_share_output_value_annual
Biaya Pekerja Produksi/Total Biaya Pekerja → ind_production_labor_cost_share_total_labor_cost_annual
Biaya Pekerja Lainnya/Total Biaya Pekerja → ind_nonproduction_labor_cost_share_total_labor_cost_annual
Upah Pekerja Produksi/Jumlah Tenaga Kerja Produksi → ind_production_wage_per_production_worker_annual
Upah Pekerja Lainnya/Jumlah Tenaga Kerja Lainnya → ind_nonproduction_wage_per_nonproduction_worker_annual
Proporsi Perusahaan PMDN → ind_firms_pmdn_share_annual
Proporsi Perusahaan PMA → ind_firms_pma_share_annual
Pertumbuhan Nilai Output → growth_ind_output_value_annual
Pertumbuhan Biaya Input → growth_ind_input_cost_annual
Pertumbuhan Nilai Tambah Pasar → growth_ind_value_added_market_annual
Pertumbuhan Nilai Tambah Produksi → growth_ind_value_added_factor_cost_annual
Pertumbuhan Tenaga Kerja → growth_ind_workers_total_annual
Pertumbuhan Jumlah Perusahaan → growth_ind_firms_total_annual
Pertumbuhan Pengeluaran Tenaga Kerja → growth_ind_labor_cost_total_annual
Pertumbuhan Produktivitas (Nilai Tambah Pasar) Tenaga Kerja → growth_ind_labor_productivity_market_annual
Pertumbuhan Produktivitas (Nilai Tambah Produksi) Tenaga Kerja → growth_ind_labor_productivity_factor_cost_annual
Selisih Pertumbuhan Biaya Input dan Pertumbuhan Nilai Output → diff_growth_input_cost_minus_output_value_annual
Selisih Pertumbuhan Biaya Pekerja dan Nilai Output → diff_growth_labor_cost_minus_output_value_annual
PDRB Riil (IDR miliar) → macro_pdrb_real_idr_billion_annual
PDRB Riil (IDR juta) → macro_pdrb_real_idr_million_annual
PDRB Manufaktur (IDR juta) → macro_pdrb_manufacturing_idr_million_annual
Manufacturing Share → macro_pdrb_manufacturing_share_pct_annual
Laju Pertumbuhan PDB Industri Manufaktur → macro_pdrb_manufacturing_growth_pct_annual
PMI Manufaktur → macro_pmi_manufaktur_national_annual
FDI → fin_fdi_annual
PDRB/Pekerja → macro_pdrb_per_worker_annual
Laju Pertumbuhan PDB Per Tenaga Kerja → macro_pdrb_per_worker_growth_pct_annual
ICOR → macro_icor_annual
Inflasi YoY (per IV) → price_inflation_yoy_q4_pct_annual
Inflasi YoY (Average) → price_inflation_yoy_avg_pct_annual
IHP (2016=100) → price_producer_index_annual
IHK (2010=100) → price_cpi_index_annual
Perubahan harga produsen (%) → price_producer_change_pct_annual
Perubahan harga konsumen (%) → price_consumer_change_pct_annual
Nilai Ekspor BPS → trade_export_value_bps_annual
Nilai Ekspor → trade_export_value_annual
Nilai Impor → trade_import_value_annual
Neraca Perdagangan → trade_balance_annual
NPL → fin_npl_annual
BI Rate (Average)* → macro_bi_rate_avg_pct_national_annual
Kurs → macro_exchange_rate_idr_usd_national_annual
IKK → macro_consumer_confidence_index_annual
"""
MAP = """
year → year
name → region_level
id → region_id
id_label → region_name
nonpkwtt → map_nonpkwtt_annual
wageump → map_wageump_annual
laborinten → map_labor_intensity_annual
formalgrowth → map_formal_growth_annual
pdrbgrowth → map_pdrb_growth_annual
pdrb → map_pdrb_annual
unemploy → map_unemployment_annual
unemploy_1529 → map_youth_unemployment_annual
informal → map_informal_annual
manufacture → map_manufacturing_annual
agriculture → map_agriculture_annual
nonmanufacture → map_nonmanufacturing_annual
services → map_services_annual
minwage → map_minimum_wage_annual
minwagegrowth → map_minimum_wage_growth_annual
threshold_skortotal → map_threshold_score_total_annual
threshold_status → map_threshold_status_annual
"""
def parse(block):
    d={}
    for line in block.strip().splitlines():
        if '→' not in line: continue
        o,f=line.split('→',1)
        d[norm(o)]=f.strip()
    return d
USER={'Database (Bulan)':parse(BULAN),'Database (Triwulan)':parse(TRIWULAN),
      'Database (Tahun)':parse(TAHUN),'MAP':parse(MAP)}
# MAP quantile_* wildcard
MAP_QUANTILE=lambda h: 'map_quantile_'+h[len('quantile_'):]+'_annual' if h.startswith('quantile_') else None

# ---- actual columns ----
def headers(path, sheet, hdr_row=1):
    wb=load_workbook(path, read_only=True, data_only=True); ws=wb[sheet]
    it=ws.iter_rows(values_only=True)
    for _ in range(hdr_row-1): next(it)
    return [norm(c) for c in next(it) if c is not None and norm(c)!='']

actual={
 'Database (Bulan)':    headers(PHKX,'Database (Bulan)'),
 'Database (Triwulan)': headers(PHKX,'Database (Triwulan)'),
 'Database (Tahun)':    headers(PHKX,'Database (Tahun)'),
 'MAP':                 headers(MAPX,'Sheet 1 - MAP_composite_all', hdr_row=2),
}

# ---- coverage check ----
print("=== COVERAGE: actual columns not in user mapping ===")
uncovered={}
for sheet, cols in actual.items():
    um=USER[sheet]
    miss=[c for c in cols if c not in um and not (sheet=='MAP' and MAP_QUANTILE(c))]
    if miss:
        uncovered[sheet]=miss
        print(f"\n[{sheet}] {len(miss)} uncovered:")
        for c in miss: print("   -", c)
if not uncovered: print("  none — every actual column is covered by the spec.")
print("\n=== user mappings with NO matching actual column (spec typo / removed) ===")
for sheet, um in USER.items():
    extra=[o for o in um if o not in actual[sheet]]
    if extra:
        print(f"[{sheet}]:")
        for o in extra: print("   -", repr(o), "->", um[o])

# ================= build the full dictionary =================
# carry over direction_risk / r_squared from the existing dictionary (join on original name)
oldmeta={}
try:
    for r in csv.DictReader(open('dictionary/indicator_dictionary.csv')):
        oldmeta[norm(r.get('original_column_name',''))]=r
except FileNotFoundError:
    pass

IDENTITY={'province_name','year','month','date','quarter','quarter_end_month_label',
          'region_name','region_id','region_level'}
def domain(fn):
    if fn in IDENTITY: return 'identity'
    for p in ['emp_share','bpjstk','phk','lab','wage','macro','price','trade','fin','emp','ind','map','geo']:
        if fn.startswith(p+'_') or fn==p: return p
    if fn.startswith('growth_'): return 'growth'
    if fn.startswith('lag1_') or fn.startswith('lag_'): return 'lag'
    if fn.startswith('diff_'): return 'diff'
    return 'manual_review_required'
def frequency(sheet, fn):
    if sheet=='Database (Bulan)': return 'monthly'
    if sheet=='Database (Triwulan)': return 'quarterly'
    if sheet in ('Database (Tahun)','MAP'): return 'annual'
    return 'n/a'
def geolevel(fn, sheet):
    if fn in ('region_name','region_id','region_level'): return 'geo_key'
    if '_national' in fn: return 'national'
    if sheet=='MAP': return 'province_or_kabkot'
    return 'province'
def unit(fn):
    for suf,u in [('_idr_billion','IDR billion'),('_idr_million','IDR million'),
                  ('_usd_million','USD million'),('_usd_per_barrel','USD per barrel'),
                  ('_idr_usd','IDR per USD'),('_pp','percentage points'),('_pct','percent'),
                  ('_index','index'),('_idr','IDR')]:
        if suf in fn: return u
    if fn.startswith(('emp_','ind_workers','ind_firms')) and fn.endswith('_annual'): return 'count'
    if fn.startswith('lab_') and any(k in fn for k in ('workers','population','count','employed','hours','seekers','vacancies','placements','underemployed')): return 'count'
    if fn.startswith('bpjstk_'): return 'count'
    return 'manual_review_required'

# national columns (empirically constant across provinces) from prior detection
NATL={'macro_pmi_manufaktur_national','macro_bi_rate_pct_national','macro_exchange_rate_idr_usd_national',
      'price_brent_oil_usd_per_barrel_national','price_ihpb_national',
      'macro_pmi_manufaktur_national_quarterly','macro_bi_rate_pct_national_quarterly',
      'macro_exchange_rate_idr_usd_national_quarterly','price_producer_index_quarterly',
      'price_producer_change_pct_quarterly','price_producer_change_ceic_pct_quarterly',
      'macro_pmi_manufaktur_national_annual','price_producer_index_annual','price_cpi_index_annual',
      'price_producer_change_pct_annual','price_consumer_change_pct_annual',
      'macro_bi_rate_avg_pct_national_annual','macro_consumer_confidence_index_annual'}
FEATURE={'Database (Bulan)':'phk_monthly_features.csv','Database (Triwulan)':'phk_quarterly_features.csv',
         'Database (Tahun)':'phk_annual_features.csv','MAP':'map_kabkot_context.csv (kabkot rows)'}

# MAP twins UNIFIED into a PHK column in the province master, then dropped.
# map_final -> (phk_final, source_flag, phk_unit_prescale)
TWIN_OF={
 'map_informal_annual':      ('lab_informal_employment_share_pct_annual','lab_informal_source',1),
 'map_unemployment_annual':  ('lab_tpt_pct_annual','lab_tpt_source',1),
 'map_minimum_wage_annual':  ('wage_ump_idr_annual','wage_ump_source',1000),
 'map_manufacturing_annual': ('emp_share_manufacturing_pct_annual','emp_share_manufacturing_source',100),
 'map_agriculture_annual':   ('emp_share_agriculture_forestry_fishery_pct_annual','emp_share_agriculture_source',100),
 'map_pdrb_annual':          ('macro_pdrb_real_idr_billion_annual','macro_pdrb_real_source',1),
}
PHK_UNIFIED={p:(m,s,pre) for m,(p,s,pre) in TWIN_OF.items()}

FIELDS=['source_dataset','source_file','source_sheet','original_column_name','final_column_name',
        'domain','indicator_label_id','indicator_label_en','definition','geographic_level','frequency',
        'unit','direction_risk','r_squared','transformation_rule','merge_key',
        'include_in_master','clean_feature_file','unified_into','national_repeated','manual_review_required','notes']
rows=[]
def add(**k):
    rows.append({f:k.get(f,'') for f in FIELDS})

SRCFILE={'Database (Bulan)':'Data untuk PHK Dashboard.xlsx','Database (Triwulan)':'Data untuk PHK Dashboard.xlsx',
         'Database (Tahun)':'Data untuk PHK Dashboard.xlsx','MAP':'MAP_composite_all.xlsx'}
SRCDS={'MAP':'map_composite'}
for sheet in ['Database (Bulan)','Database (Triwulan)','Database (Tahun)','MAP']:
    um=USER[sheet]
    for h in actual[sheet]:
        fn=um.get(h) or (MAP_QUANTILE(h) if sheet=='MAP' else None)
        mr='yes' if fn is None else 'no'
        if fn is None: fn='manual_review_required'
        dom=domain(fn); freq=frequency(sheet,fn); geo=geolevel(fn,sheet); un=unit(fn)
        isident= fn in IDENTITY
        isnatl= fn in NATL
        unified_into=''
        note=''
        if isident:
            inc='key'; feat=''
        elif fn in TWIN_OF:                       # MAP twin -> unified into a PHK column, then dropped
            phk_c,src,_pre=TWIN_OF[fn]; phk_c=shorten(phk_c); src=shorten(src)
            inc='no'; feat='map_kabkot_context.csv (kabkot rows)'; unified_into=phk_c
            note=f'unified into {phk_c} in the province master (MAP backfills 2014-2021); {src} flags origin'
        elif sheet=='MAP':
            inc='yes (province rows only)'; feat='map_kabkot_context.csv (kabkot rows)'
        elif isnatl:
            inc='no'; feat=f'national_{freq}.csv'
        else:
            inc='yes'; feat=FEATURE[sheet]
        transf=''
        if not isident and not isnatl and sheet=='Database (Tahun)': transf='repeat annual value across 12 months (flag_annual_repeated_monthly)'
        if not isident and not isnatl and sheet=='Database (Triwulan)': transf='repeat quarterly value across 3 months (flag_quarterly_repeated_monthly)'
        if sheet=='MAP' and fn in TWIN_OF: transf='annual; repeat across months; consumed by unification'
        elif sheet=='MAP' and not isident: transf='annual; repeat across months; 2024 carried into 2025'
        if fn in PHK_UNIFIED:                     # PHK side of a unified pair
            mapc,src,pre=PHK_UNIFIED[fn]; mapc=shorten(mapc); src=shorten(src)
            if pre!=1: transf=(transf+'; ' if transf else '')+f'unit x{pre} (harmonize to {("percent" if pre==100 else "IDR")})'
            note=f'UNIFIED: PHK where available + {mapc} backfill for 2014-2021; {src} flags origin per row'
        om=oldmeta.get(h,{})
        add(source_dataset=SRCDS.get(sheet,'phk_dashboard'), source_file=SRCFILE[sheet], source_sheet=sheet,
            original_column_name=h, final_column_name=shorten(fn), domain=dom, indicator_label_id=h,
            indicator_label_en=(om.get('indicator_label_en','') if om.get('indicator_label_en','') not in ('','manual_review_required') else ''),
            definition=(om.get('definition','') if om.get('definition','') not in ('','manual_review_required') else 'manual_review_required'),
            geographic_level=geo, frequency=freq, unit=un,
            direction_risk=om.get('direction_risk','') or '', r_squared=om.get('r_squared','') or '',
            transformation_rule=transf, merge_key=('province_name_std + year + month' if sheet=='Database (Bulan)' else 'province_name_std + year + quarter' if sheet=='Database (Triwulan)' else 'province_name_std + year' if sheet=='Database (Tahun)' else 'region + year'),
            include_in_master=inc, clean_feature_file=feat, unified_into=unified_into,
            national_repeated=('yes' if isnatl else ''),
            manual_review_required=('yes' if (mr=='yes' or un=='manual_review_required' or dom=='manual_review_required') else 'no'),
            notes=note or ('unit needs confirmation' if un=='manual_review_required' else ''))
# excluded sheets
add(source_dataset='phk_dashboard', source_file='Data untuk PHK Dashboard.xlsx', source_sheet='phk',
    original_column_name='Stock / Flow (wide province x date matrix)', final_column_name='excluded_from_master',
    domain='phk', geographic_level='province', frequency='monthly', unit='count',
    include_in_master='no', clean_feature_file='(raw staging)', manual_review_required='no',
    definition='raw wide pivot of layoff Stock/Flow; tidy version already in Database (Bulan) as phk_stock / phk_flow',
    notes='excluded_from_master: redundant raw layout, not tidy')
for c in headers(PHKX,'Kebutuhan Data'):
    add(source_dataset='phk_dashboard', source_file='Data untuk PHK Dashboard.xlsx', source_sheet='Kebutuhan Data',
        original_column_name=c, final_column_name='excluded_from_master', domain='metadata',
        geographic_level='n/a', frequency='n/a', unit='n/a', include_in_master='no',
        clean_feature_file='indicator_dictionary.csv', manual_review_required='no',
        definition='requirements/metadata sheet (not observational data); used to populate the data dictionary',
        notes='excluded_from_master: metadata sheet')

with open('data/processed/indicator_dictionary.csv','w',newline='') as f:
    w=csv.DictWriter(f, fieldnames=FIELDS); w.writeheader(); w.writerows(rows)

# unify pairs (shortened names) for gen_pipeline.py: phk_col, map_col, source_flag, phk_unit_prescale
with open('data/processed/_unify_pairs.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['phk_col','map_col','source_flag','phk_prescale'])
    for mapc,(phkc,src,pre) in TWIN_OF.items():
        w.writerow([shorten(phkc), shorten(mapc), shorten(src), pre])

# fail loudly if any final name still exceeds Stata's 32-char limit
over=sorted({r['final_column_name'] for r in rows if len(r['final_column_name'])>32})
assert not over, f"names >32 chars: {over}"
print(f"max final-name length: {max(len(r['final_column_name']) for r in rows)} (<=32 OK)")

# ---- summary ----
from collections import Counter
print("\n=== dictionary written: data/processed/indicator_dictionary.csv ===")
print("total rows:", len(rows))
inc=Counter(r['include_in_master'] for r in rows); print("include_in_master:", dict(inc))
print("by domain:", dict(Counter(r['domain'] for r in rows)))
print("manual_review_required=yes:", sum(1 for r in rows if r['manual_review_required']=='yes'))
print("national:", sum(1 for r in rows if r['national_repeated']=='yes'))

